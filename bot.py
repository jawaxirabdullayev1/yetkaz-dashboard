import asyncio
import html
import io
import json
import logging
import re
import sqlite3
from collections import Counter
from datetime import datetime
from pathlib import Path

from aiogram import Bot, Dispatcher, F
from aiogram.client.default import DefaultBotProperties
from aiogram.filters import Command, CommandObject, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.base import BaseStorage, StorageKey
from aiogram.types import (
    BufferedInputFile, CallbackQuery, InlineKeyboardButton, InlineKeyboardMarkup,
    KeyboardButton, Message, ReplyKeyboardMarkup, ReplyKeyboardRemove,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import branches
import config
from regions import REGIONS

BASE = Path(__file__).parent
DB = BASE / config.DB_PATH

# ===================== BAZA =====================

def db():
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    return con


def init_db():
    with db() as con:
        con.execute("""CREATE TABLE IF NOT EXISTS leads(
            user_id INTEGER PRIMARY KEY,
            username TEXT, name TEXT, phone TEXT,
            source TEXT, answers TEXT, score INTEGER, level TEXT,
            created TEXT, updated TEXT)""")
        cols = {r["name"] for r in con.execute("PRAGMA table_info(leads)")}
        for c in ("viloyat", "tuman", "lat", "lon", "maps", "holat",
                  "xulosa", "sabab", "eng_yaqin", "masofa_km", "tumanda_filial", "zichlik", "yaqinlar",
                  "bosqich", "answers2", "photos", "status", "izoh", "aloqa_sana", "status_vaqt", "filial_id", "korildi"):
            if c not in cols:
                con.execute(f"ALTER TABLE leads ADD COLUMN {c} TEXT")
        con.execute("CREATE TABLE IF NOT EXISTS settings(k TEXT PRIMARY KEY, v TEXT)")
        con.execute("CREATE TABLE IF NOT EXISTS fsm(k TEXT PRIMARY KEY, state TEXT, data TEXT)")


class SQLiteStorage(BaseStorage):
    """Anketa holatini bazada saqlaydi — bot qayta ishga tushsa ham anketa davom etadi."""

    @staticmethod
    def _k(key: StorageKey):
        return f"{key.bot_id}:{key.chat_id}:{key.user_id}"

    def _row(self, key):
        with db() as con:
            return con.execute("SELECT state, data FROM fsm WHERE k=?", (self._k(key),)).fetchone()

    def _upsert(self, key, field, value):
        with db() as con:
            con.execute("INSERT OR IGNORE INTO fsm(k, data) VALUES(?, '{}')", (self._k(key),))
            con.execute(f"UPDATE fsm SET {field}=? WHERE k=?", (value, self._k(key)))

    async def set_state(self, key, state=None):
        self._upsert(key, "state", getattr(state, "state", state))

    async def get_state(self, key):
        r = self._row(key)
        return r["state"] if r else None

    async def set_data(self, key, data):
        self._upsert(key, "data", json.dumps(data, ensure_ascii=False))

    async def get_data(self, key):
        r = self._row(key)
        return json.loads(r["data"] or "{}") if r else {}

    async def close(self):
        pass


def get_setting(k):
    with db() as con:
        r = con.execute("SELECT v FROM settings WHERE k=?", (k,)).fetchone()
        return r["v"] if r else None


def set_setting(k, v):
    with db() as con:
        con.execute("INSERT OR REPLACE INTO settings VALUES(?,?)", (k, str(v)))


def save_lead(d):
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    with db() as con:
        old = con.execute("SELECT created FROM leads WHERE user_id=?", (d["user_id"],)).fetchone()
        row = {**d, "answers": json.dumps(d["answers"], ensure_ascii=False),
               "created": old["created"] if old else now, "updated": now}
        keys = list(row)
        # Upsert: berilmagan ustunlar (xulosa, bosqich, ...) saqlanib qoladi
        updates = ", ".join(f"{k}=excluded.{k}" for k in keys if k not in ("user_id", "created"))
        con.execute(f"INSERT INTO leads({','.join(keys)}) VALUES({','.join('?' * len(keys))}) "
                    f"ON CONFLICT(user_id) DO UPDATE SET {updates}", [row[k] for k in keys])
    return old is None


def set_fields(user_id, **kw):
    with db() as con:
        con.execute(f"UPDATE leads SET {', '.join(k + '=?' for k in kw)} WHERE user_id=?",
                    [*kw.values(), user_id])


def get_lead(user_id):
    return next((l for l in all_leads() if l["user_id"] == user_id), None)


DONE = "✅ To'liq"
DRAFT = "⏳ Tugallanmagan"


def save_draft(chat, data):
    """Har bir javobdan keyin chala anketani saqlaydi (to'liq anketani buzmaydi)."""
    if not data.get("name"):
        return
    with db() as con:
        old = con.execute("SELECT holat FROM leads WHERE user_id=?", (chat.id,)).fetchone()
    if old and (old["holat"] or DONE) == DONE:
        return
    save_lead(lead_from(chat, data, DRAFT))


def lead_from(chat, d, holat):
    return {
        "user_id": chat.id, "username": chat.username or "", "name": html.escape(d.get("name", "")),
        "phone": d.get("phone", ""), "viloyat": d.get("viloyat"), "tuman": d.get("tuman"),
        "lat": d.get("lat"), "lon": d.get("lon"), "maps": d.get("maps"), "source": d.get("source", ""),
        "answers": d.get("answers", {}), "score": d.get("score", 0),
        "level": level_for(d.get("score", 0)), "holat": holat,
    }


async def step(state, chat, **kw):
    await state.update_data(**kw)
    save_draft(chat, await state.get_data())


def all_leads():
    with db() as con:
        rows = con.execute("SELECT * FROM leads ORDER BY (holat = ?), "
                           "CASE WHEN xulosa LIKE '%Ma''qul emas%' THEN 2 WHEN xulosa LIKE '%Ma''qul%' THEN 0 "
                           "ELSE 1 END, score DESC, created DESC", (DRAFT,)).fetchall()
    out = []
    for r in rows:
        d = dict(r)
        d["answers"] = json.loads(d["answers"] or "{}")
        d["viloyat"] = d["viloyat"] or "-"
        d["tuman"] = d["tuman"] or "-"
        d["holat"] = d["holat"] or DONE
        d["xulosa"] = d["xulosa"] or ("—" if d["holat"] == DRAFT else "⏳ Tekshirilmagan")
        d["yaqinlar"] = json.loads(d["yaqinlar"] or "[]")
        d["answers2"] = json.loads(d["answers2"] or "{}")
        d["photos"] = json.loads(d["photos"] or "{}")
        d["bosqich"] = d["bosqich"] or ""
        d["status"] = d["status"] or "Yangi"
        d["izoh"] = d["izoh"] or ""
        d["aloqa_sana"] = d["aloqa_sana"] or ""
        out.append(d)
    return out

# ===================== YORDAMCHI =====================

def level_for(score):
    if score >= config.HOT_SCORE:
        return "🔥 Issiq"
    if score >= config.WARM_SCORE:
        return "🌤 Iliq"
    return "❄️ Sovuq"


def map_link(l):
    if l.get("lat"):
        return f"https://maps.google.com/?q={l['lat']},{l['lon']}"
    return l.get("maps") or ""


def qlabel(q):
    """Savolning birinchi qatori (izohsiz) — hisobot va Excel uchun."""
    return q["text"].split("\n")[0]


def _norm(text):
    text = text.lower()
    for ch in "'‘’ʻʼ`":
        text = text.replace(ch, "")
    return re.sub(r"[^\w]+", " ", text).strip()


NO_ANSWERS = {_norm(w) for w in (
    "yo'q", "yoq", "yuq", "yo'q hali", "hali yo'q", "hozircha yo'q", "hech qanday",
    "нет", "net", "no", "-", "0")}


def is_negative(text):
    return _norm(text) in NO_ANSWERS or not _norm(text)


def is_admin(uid):
    return uid in config.ADMIN_IDS

# ===================== ANKETA =====================

class Form(StatesGroup):
    name = State()
    location = State()
    viloyat = State()
    tuman = State()
    phone = State()
    question = State()


dp = Dispatcher(storage=SQLiteStorage())
VILOYATLAR = list(REGIONS)
MAPS_RE = re.compile(
    r"https?://\S*(google\.[a-z.]+/maps|maps\.google\.|maps\.app\.goo\.gl|goo\.gl/maps|yandex\.\S+/maps)\S*",
    re.I)


def grid(buttons, cols=2):
    return [buttons[i:i + cols] for i in range(0, len(buttons), cols)]


def viloyat_kb():
    return InlineKeyboardMarkup(inline_keyboard=grid(
        [InlineKeyboardButton(text=v, callback_data=f"v:{i}") for i, v in enumerate(VILOYATLAR)]))


def tuman_kb(vi):
    rows = grid([InlineKeyboardButton(text=t, callback_data=f"t:{vi}:{ti}")
                 for ti, t in enumerate(REGIONS[VILOYATLAR[vi]])])
    rows.append([InlineKeyboardButton(text="⬅️ Viloyatni o'zgartirish", callback_data="t:back")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def question_kb(qi):
    opts = config.QUESTIONS[qi]["options"]
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=t, callback_data=f"q:{qi}:{oi}")]
        for oi, (t, _) in enumerate(opts)
    ])


@dp.message(CommandStart())
async def start(msg: Message, command: CommandObject, state: FSMContext):
    await state.clear()
    source = (command.args or "organik")[:64]
    await state.update_data(source=source, answers={}, score=0)
    await msg.answer(config.GREETING + "\n\nIsm va familiyangizni kiriting:",
                     reply_markup=ReplyKeyboardRemove())
    await state.set_state(Form.name)


@dp.message(Form.name, F.text)
async def got_name(msg: Message, state: FSMContext):
    name = msg.text.strip()
    if len(name) < 2 or name.startswith("/"):
        return await msg.answer("Iltimos, ism va familiyangizni to'g'ri yozing:")
    await step(state, msg.chat, name=name[:100])
    kb = ReplyKeyboardMarkup(resize_keyboard=True, one_time_keyboard=True, keyboard=[
        [KeyboardButton(text="📍 Hozirgi joylashuvimni yuborish", request_location=True)]])
    await msg.answer(
        "Filial ochiladigan joyning lokatsiyasini yuboring:\n\n"
        "• hozir o'sha joyda bo'lsangiz — pastdagi tugmani bosing;\n"
        "• boshqa joyda bo'lsangiz — 📎 → <b>Joylashuv</b> orqali xaritadan nuqtani tanlang "
        "yoki Google Maps havolasini yuboring.", reply_markup=kb)
    await state.set_state(Form.location)


@dp.message(Form.location, F.location)
async def got_location(msg: Message, state: FSMContext):
    await step(state, msg.chat, lat=str(msg.location.latitude), lon=str(msg.location.longitude), maps=None)
    await _ask_viloyat(msg, state, "Lokatsiya qabul qilindi ✅")


@dp.message(Form.location, F.text.regexp(MAPS_RE))
async def got_maps_link(msg: Message, state: FSMContext):
    await step(state, msg.chat, lat=None, lon=None, maps=MAPS_RE.search(msg.text).group(0)[:300])
    await _ask_viloyat(msg, state, "Havola qabul qilindi ✅")


@dp.message(Form.location)
async def location_other(msg: Message):
    await msg.answer("Iltimos, lokatsiya yuboring (tugma yoki 📎 → Joylashuv) "
                     "yoki Google Maps havolasini tashlang.")


async def _ask_viloyat(msg, state, prefix):
    await msg.answer(prefix, reply_markup=ReplyKeyboardRemove())
    await msg.answer("Faoliyat yuritadigan viloyatingizni tanlang:", reply_markup=viloyat_kb())
    await state.set_state(Form.viloyat)


@dp.callback_query(Form.viloyat, F.data.startswith("v:"))
async def got_viloyat(cb: CallbackQuery, state: FSMContext):
    vi = int(cb.data[2:])
    await step(state, cb.message.chat, viloyat=VILOYATLAR[vi])
    await cb.message.edit_text(f"Viloyat: <b>{VILOYATLAR[vi]}</b>\n\nTuman (shahar)ni tanlang:",
                               reply_markup=tuman_kb(vi))
    await state.set_state(Form.tuman)
    await cb.answer()


@dp.callback_query(Form.tuman, F.data == "t:back")
async def tuman_back(cb: CallbackQuery, state: FSMContext):
    await cb.message.edit_text("Faoliyat yuritadigan viloyatingizni tanlang:", reply_markup=viloyat_kb())
    await state.set_state(Form.viloyat)
    await cb.answer()


@dp.callback_query(Form.tuman, F.data.startswith("t:"))
async def got_tuman(cb: CallbackQuery, state: FSMContext):
    _, vi, ti = cb.data.split(":")
    viloyat = VILOYATLAR[int(vi)]
    tuman = REGIONS[viloyat][int(ti)]
    await step(state, cb.message.chat, viloyat=viloyat, tuman=tuman)
    await cb.message.edit_text(f"📍 <b>{viloyat}, {tuman}</b>")
    await cb.answer()
    kb = ReplyKeyboardMarkup(resize_keyboard=True, one_time_keyboard=True, keyboard=[
        [KeyboardButton(text="📱 Raqamni yuborish", request_contact=True)]])
    await cb.message.answer("Siz bilan bog'lanish uchun telefon raqamingizni yuboring "
                            "(tugmani bosing yoki yozing):", reply_markup=kb)
    await state.set_state(Form.phone)


@dp.message(Form.phone, F.contact)
async def got_contact(msg: Message, state: FSMContext):
    await _set_phone(msg, state, msg.contact.phone_number)


@dp.message(Form.phone, F.text)
async def got_phone_text(msg: Message, state: FSMContext):
    digits = re.sub(r"\D", "", msg.text)
    if len(digits) == 9:
        digits = "998" + digits
    if not (len(digits) == 12 and digits.startswith("998")):
        return await msg.answer("Raqam noto'g'ri. Masalan: +998 90 123 45 67")
    await _set_phone(msg, state, digits)


async def _set_phone(msg, state, phone):
    phone = "+" + re.sub(r"\D", "", phone)
    await step(state, msg.chat, phone=phone)
    await msg.answer("Rahmat ✅", reply_markup=ReplyKeyboardRemove())
    await ask_question(msg, state, 0)


async def ask_question(msg, state, qi):
    if qi >= len(config.QUESTIONS):
        return await finish(msg, state)
    await state.update_data(qi=qi)
    await state.set_state(Form.question)
    q = config.QUESTIONS[qi]
    num = f"<b>{qi + 1}/{len(config.QUESTIONS)}.</b> "
    if q.get("type") == "choice":
        await msg.answer(num + q["text"], reply_markup=question_kb(qi))
    else:
        await msg.answer(num + q["text"])


@dp.message(Form.question, F.text)
async def got_text_answer(msg: Message, state: FSMContext):
    data = await state.get_data()
    qi = data["qi"]
    q = config.QUESTIONS[qi]
    if q.get("type") == "choice":
        return await msg.answer("Iltimos, yuqoridagi tugmalardan birini tanlang 👆")
    text = msg.text.strip()
    if not text or text.startswith("/"):
        return await msg.answer("Iltimos, javobingizni yozing:")
    answers = data["answers"]
    answers[q["key"]] = html.escape(text[:1000])
    pts = 0 if is_negative(text) else q.get("points", 0)
    await step(state, msg.chat, answers=answers, score=data["score"] + pts)
    await ask_question(msg, state, qi + 1)


@dp.callback_query(Form.question, F.data.startswith("q:"))
async def got_choice(cb: CallbackQuery, state: FSMContext):
    _, qi, oi = cb.data.split(":")
    qi, oi = int(qi), int(oi)
    data = await state.get_data()
    if qi != data.get("qi"):
        return await cb.answer()
    q = config.QUESTIONS[qi]
    text, pts = q["options"][oi]
    answers = data["answers"]
    answers[q["key"]] = text
    await step(state, cb.message.chat, answers=answers, score=data["score"] + pts)
    await cb.message.edit_text(f"{qlabel(q)}\n<b>{text}</b>")
    await cb.answer()
    await ask_question(cb.message, state, qi + 1)


async def finish(msg: Message, state: FSMContext):
    d = await state.get_data()
    await state.clear()
    with db() as con:
        old = con.execute("SELECT holat, bosqich FROM leads WHERE user_id=?", (msg.chat.id,)).fetchone()
    is_new = not old or old["holat"] == DRAFT
    lead = lead_from(msg.chat, d, DONE)
    save_lead(lead)
    set_fields(msg.chat.id, korildi="")   # yangi / yangilangan ariza dashboardda "YANGI" bo'lib ko'rinadi
    if old and old["bosqich"]:
        # Anketa qayta to'ldirildi — taklif bosqichi boshidan (eski javoblar keyingi to'ldirishda almashadi)
        set_fields(msg.chat.id, bosqich="")
        lead["oldingi_bosqich"] = old["bosqich"]
    await msg.answer(config.FINISH_TEXT)
    lead.update(await run_evaluation(lead))
    await notify_admins(msg.bot, lead, is_new)


async def run_evaluation(lead):
    """Mavjud filiallar bilan solishtirib, natijani bazaga yozadi."""
    try:
        res = await asyncio.to_thread(branches.evaluate, lead)
    except Exception as e:
        logging.exception("Tahlil xatosi")
        res = {"xulosa": branches.KORIB, "sabab": f"Avtomatik tahlil ishlamadi: {e}"}
    res.setdefault("yaqinlar", [])
    cols = ["xulosa", "sabab", "eng_yaqin", "masofa_km", "tumanda_filial", "zichlik", "lat", "lon"]
    sets = ", ".join(f"{c}=COALESCE(?, {c})" for c in cols)
    vals = [res.get(c) for c in cols] + [json.dumps(res["yaqinlar"], ensure_ascii=False), lead["user_id"]]
    with db() as con:
        con.execute(f"UPDATE leads SET {sets}, yaqinlar=? WHERE user_id=?", vals)
    return {k: v for k, v in res.items() if v is not None}


async def notify_admins(bot: Bot, lead, is_new):
    ans = "\n".join(f"• {qlabel(q)}\n   <b>{lead['answers'].get(q['key'], '-')}</b>"
                    for q in config.QUESTIONS)
    tags = " ".join("#" + re.sub(r"\W+", "_", x).strip("_") for x in (lead["viloyat"], lead["tuman"]))
    link = map_link(lead)
    geo = f"\n🗺 <a href=\"{html.escape(link)}\">Xaritada ko'rish</a>" if link else ""
    verdict = ""
    if lead.get("xulosa"):
        near = "\n".join(f"   • {x}" for x in lead.get("yaqinlar") or [])
        verdict = (f"<b>XULOSA: {lead['xulosa']}</b>\n{html.escape(lead.get('sabab') or '', quote=False)}\n"
                   f"🏢 Tumanda aktiv filial: {lead.get('tumanda_filial', '-')}\n"
                   f"👥 Aholi: {lead.get('zichlik') or '-'}\n"
                   + (f"📏 Eng yaqin filiallar:\n{near}\n" if near else "") + "\n")
    text = (verdict + f"{'🆕 Yangi lid' if is_new else '♻️ Lid yangilandi'} — {lead['level']} ({lead['score']} ball)\n\n"
            f"👤 {lead['name']}" + (f" (@{lead['username']})" if lead["username"] else "") +
            f"\n📞 {lead['phone']}\n📍 {lead['viloyat']}, {lead['tuman']}{geo}\n"
            f"🔗 Manba: {lead['source']}\n\n{ans}\n\n{tags} #{lead['source']}")
    if lead.get("oldingi_bosqich"):
        text += (f"\n\n⚠️ Bu odam anketani qayta to'ldirdi. Oldingi taklif bosqichi: {lead['oldingi_bosqich']}. "
                 f"Taklif jarayoni boshidan boshlanadi.")
    markup = None
    if (lead.get("xulosa") or "")[:1] in ("✅", "⚠"):
        markup = InlineKeyboardMarkup(inline_keyboard=[[
            InlineKeyboardButton(text="📨 Davom ettirish (taklif yuborish)", callback_data=f"go:{lead['user_id']}")]])
    await send_admins(bot, text, markup)


def admin_targets():
    group = get_setting("group")
    return {int(group)} if group else set(config.ADMIN_IDS)


async def send_admins(bot: Bot, text, markup=None):
    for t in admin_targets():
        try:
            await bot.send_message(t, text[:4096], disable_web_page_preview=True, reply_markup=markup)
        except Exception as e:
            logging.warning("Xabar yuborilmadi %s: %s", t, e)

# ===================== IKKINCHI BOSQICH: TAKLIF VA QO'SHIMCHA MA'LUMOT =====================

OFFER_SENT = "📨 Taklif yuborildi"
ASKED = "❓ Savol berdi"
AGREED = "🤝 Roziman (ma'lumot to'ldirmoqda)"
REFUSED = "🚫 Rad etdi"
COMPLETE = "📋 Ma'lumot to'liq"
APPROVED = "✅ Hujjatlar tasdiqlandi"


class Stage2(StatesGroup):
    offer = State()
    ask = State()
    question = State()


def offer_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ Roziman", callback_data="of:yes")],
        [InlineKeyboardButton(text="❓ Savolim bor", callback_data="of:ask"),
         InlineKeyboardButton(text="❌ Rad etaman", callback_data="of:no")]])


def user_state(bot: Bot, uid):
    return FSMContext(storage=dp.storage, key=StorageKey(bot_id=bot.id, chat_id=uid, user_id=uid))


async def send_offer(bot: Bot, lead):
    """Lidga tijoriy taklifni yuboradi. Muvaffaqiyatli bo'lsa True."""
    text = (BASE / config.OFFER_FILE).read_text(encoding="utf-8")
    name = html.unescape(lead["name"]).split()[0] if lead["name"] else ""
    try:
        await bot.send_message(lead["user_id"], f"Assalomu alaykum, {name}! 👋\n"
                               f"Arizangiz ko'rib chiqildi va hududingiz bizga ma'qul keldi.",
                               parse_mode=None)
        await bot.send_message(lead["user_id"], text, parse_mode=None, reply_markup=offer_kb())
    except Exception as e:
        logging.warning("Taklif yuborilmadi %s: %s", lead["user_id"], e)
        return False
    st = user_state(bot, lead["user_id"])
    await st.clear()
    await st.set_state(Stage2.offer)
    set_fields(lead["user_id"], bosqich=OFFER_SENT)
    return True


@dp.callback_query(F.data.startswith("go:"))
async def admin_go(cb: CallbackQuery, bot: Bot):
    if not is_admin(cb.from_user.id):
        return await cb.answer("Faqat adminlar uchun", show_alert=True)
    lead = get_lead(int(cb.data[3:]))
    if not lead:
        return await cb.answer("Lid topilmadi", show_alert=True)
    if lead["bosqich"]:
        return await cb.answer(f"Allaqachon: {lead['bosqich']}", show_alert=True)
    ok = await send_offer(bot, lead)
    await cb.answer("Taklif yuborildi ✅" if ok else "Yuborib bo'lmadi (botni bloklagan bo'lishi mumkin)",
                    show_alert=not ok)
    if ok:
        await cb.message.edit_reply_markup(reply_markup=InlineKeyboardMarkup(inline_keyboard=[[
            InlineKeyboardButton(text=f"✅ Taklif yuborildi ({cb.from_user.first_name})", callback_data="noop")]]))


@dp.callback_query(F.data == "goall")
async def admin_go_all(cb: CallbackQuery, bot: Bot):
    if not is_admin(cb.from_user.id):
        return await cb.answer("Faqat adminlar uchun", show_alert=True)
    leads = [l for l in all_leads() if l["holat"] == DONE and l["xulosa"].startswith("✅") and not l["bosqich"]]
    await cb.answer()
    ok = 0
    for l in leads:
        ok += await send_offer(bot, l)
        await asyncio.sleep(0.5)
    await cb.message.edit_reply_markup(reply_markup=None)
    await cb.message.answer(f"📨 Taklif yuborildi: <b>{ok}</b> ta" +
                            (f", yuborilmadi: {len(leads) - ok} ta" if ok < len(leads) else ""))


@dp.callback_query(F.data == "noop")
async def noop(cb: CallbackQuery):
    await cb.answer()


def lead_title(l):
    return f"{l['name']} ({l['viloyat']}, {l['tuman']}) 📞 {l['phone']}"


@dp.callback_query(Stage2.offer, F.data == "of:yes")
async def offer_yes(cb: CallbackQuery, state: FSMContext):
    await cb.message.edit_reply_markup(reply_markup=None)
    await cb.answer()
    await agreed(cb.message, state, cb.bot)


@dp.message(Stage2.offer, F.text.lower().contains("rozi"))
async def offer_yes_text(msg: Message, state: FSMContext):
    await agreed(msg, state, msg.bot)


async def agreed(msg: Message, state: FSMContext, bot: Bot):
    set_fields(msg.chat.id, bosqich=AGREED)
    lead = get_lead(msg.chat.id)
    await send_admins(bot, f"🤝 <b>Roziman</b> dedi: {lead_title(lead)}\nQo'shimcha ma'lumot to'ldirmoqda...")
    await msg.answer("Ajoyib! 🤝 Rasmiylashtirish uchun bir nechta qo'shimcha savolga javob bering.")
    await state.update_data(answers2={}, photos={})
    await ask_followup(msg, state, 0)


@dp.callback_query(Stage2.offer, F.data == "of:no")
async def offer_no(cb: CallbackQuery, state: FSMContext):
    await cb.message.edit_reply_markup(reply_markup=None)
    await cb.answer()
    await state.clear()
    set_fields(cb.message.chat.id, bosqich=REFUSED)
    await cb.message.answer("Tushunarli, javobingiz uchun rahmat! Fikringiz o'zgarsa, bizga yozing.")
    await send_admins(cb.bot, f"🚫 <b>Rad etdi</b>: {lead_title(get_lead(cb.message.chat.id))}")


@dp.callback_query(Stage2.offer, F.data == "of:ask")
async def offer_ask(cb: CallbackQuery, state: FSMContext):
    await cb.answer()
    await state.set_state(Stage2.ask)
    await cb.message.answer("Savolingizni yozing — mutaxassisimizga yetkazamiz:")


@dp.message(Stage2.ask, F.text)
async def got_offer_question(msg: Message, state: FSMContext):
    set_fields(msg.chat.id, bosqich=ASKED)
    await send_admins(msg.bot, f"❓ <b>Savol</b> — {lead_title(get_lead(msg.chat.id))}\n\n"
                               f"{html.escape(msg.text[:3000])}")
    await state.set_state(Stage2.offer)
    await msg.answer("Savolingiz yuborildi ✅ Tez orada javob beramiz.\n\n"
                     "Taklif ma'qul bo'lsa, «✅ Roziman» tugmasini bosing yoki «Roziman» deb yozing.",
                     reply_markup=offer_kb())


@dp.message(Stage2.offer)
async def offer_other(msg: Message):
    await msg.answer("Taklif ma'qul bo'lsa «Roziman» deb yozing yoki quyidagi tugmalardan birini bosing:",
                     reply_markup=offer_kb())


async def ask_followup(msg, state, qi):
    qs = config.FOLLOWUP_QUESTIONS
    if qi >= len(qs):
        return await finish_followup(msg, state)
    await state.update_data(fqi=qi)
    await state.set_state(Stage2.question)
    q = qs[qi]
    text = f"<b>{qi + 1}/{len(qs)}.</b> {q['text']}"
    if q["type"] == "choice":
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text=o, callback_data=f"f:{qi}:{oi}")] for oi, o in enumerate(q["options"])])
    else:
        kb = None
    await msg.answer(text, reply_markup=kb)


async def save_followup(msg, state, qi, value=None, photo=None):
    data = await state.get_data()
    q = config.FOLLOWUP_QUESTIONS[qi]
    answers2, photos = data.get("answers2", {}), data.get("photos", {})
    answers2[q["key"]] = value if value is not None else "📷 rasm"
    if photo:
        photos[q["key"]] = photo
    await state.update_data(answers2=answers2, photos=photos)
    set_fields(msg.chat.id, answers2=json.dumps(answers2, ensure_ascii=False),
               photos=json.dumps(photos, ensure_ascii=False))
    if photo:
        await download_photo(msg.bot, msg.chat.id, q["key"], photo)
    await ask_followup(msg, state, qi + 1)


PHOTO_DIR = BASE / "rasmlar"


def photo_path(user_id, key):
    return PHOTO_DIR / str(user_id) / f"{key}.jpg"


async def download_photo(bot: Bot, user_id, key, file_id):
    """Rasmni kompyuterga (lid_bot/rasmlar/<user_id>/<key>.jpg) saqlaydi."""
    path = photo_path(user_id, key)
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        await bot.download(file_id[4:] if file_id.startswith("doc:") else file_id, destination=path)
        return True
    except Exception as e:
        logging.warning("Rasm saqlanmadi %s/%s: %s", user_id, key, e)
        return False


async def ensure_photos(bot: Bot, leads):
    """Hali kompyuterga tushmagan rasmlarni yuklab oladi (masalan, eski lidlar uchun)."""
    for l in leads:
        for key, fid in l["photos"].items():
            if not photo_path(l["user_id"], key).exists():
                await download_photo(bot, l["user_id"], key, fid)


@dp.message(Stage2.question, F.photo)
async def followup_photo(msg: Message, state: FSMContext):
    qi = (await state.get_data())["fqi"]
    if config.FOLLOWUP_QUESTIONS[qi]["type"] != "photo":
        return await msg.answer("Iltimos, savolga javob bering 👆")
    await save_followup(msg, state, qi, photo=msg.photo[-1].file_id)


@dp.message(Stage2.question, F.document.mime_type.startswith("image/"))
async def followup_photo_file(msg: Message, state: FSMContext):
    """Rasm fayl ko'rinishida yuborilsa ham qabul qilinadi."""
    qi = (await state.get_data())["fqi"]
    if config.FOLLOWUP_QUESTIONS[qi]["type"] != "photo":
        return await msg.answer("Iltimos, savolga javob bering 👆")
    await save_followup(msg, state, qi, photo="doc:" + msg.document.file_id)


@dp.message(Stage2.question, F.text)
async def followup_text(msg: Message, state: FSMContext):
    qi = (await state.get_data())["fqi"]
    q = config.FOLLOWUP_QUESTIONS[qi]
    if q["type"] == "photo":
        return await msg.answer("Bu savolga faqat rasm qabul qilinadi 📷\n"
                                "Rasmni yubormaguningizcha keyingi savolga o'tilmaydi.")
    if q["type"] == "choice":
        return await msg.answer("Iltimos, tugmalardan birini tanlang 👆")
    if msg.text.startswith("/"):
        return
    text = msg.text.strip()
    if q.get("require") == "phone":
        digits = re.sub(r"\D", "", text)
        if len(digits) == 9:
            digits = "998" + digits
        if not (len(digits) == 12 and digits.startswith("998")):
            return await msg.answer("Raqam noto'g'ri. Masalan: +998 90 123 45 67")
        text = "+" + digits
    if q.get("require") == "card" and not re.search(r"(?<!\d)(\d[ -]?){15}\d(?!\d)", text):
        return await msg.answer("Karta raqami topilmadi. Iltimos, 16 xonali karta raqamini ham yozing "
                                "(masalan: 8600 1234 5678 9012).")
    await save_followup(msg, state, qi, value=html.escape(text[:1000]))


@dp.message(Stage2.question)
async def followup_other(msg: Message, state: FSMContext):
    """Stiker, video, ovozli xabar va h.k. — qabul qilinmaydi."""
    qi = (await state.get_data())["fqi"]
    if config.FOLLOWUP_QUESTIONS[qi]["type"] == "photo":
        await msg.answer("Iltimos, aynan rasm yuboring 📷 (video yoki stiker emas).")
    else:
        await msg.answer("Iltimos, savolga javob bering 👆")


@dp.callback_query(Stage2.question, F.data.startswith("f:"))
async def followup_choice(cb: CallbackQuery, state: FSMContext):
    _, qi, oi = cb.data.split(":")
    qi = int(qi)
    if qi != (await state.get_data()).get("fqi"):
        return await cb.answer()
    q = config.FOLLOWUP_QUESTIONS[qi]
    value = "— (keyinroq)" if oi == "skip" else q["options"][int(oi)]
    await cb.message.edit_text(f"{q['text']}\n<b>{value}</b>")
    await cb.answer()
    await save_followup(cb.message, state, qi, value=value)


async def finish_followup(msg: Message, state: FSMContext):
    await state.clear()
    set_fields(msg.chat.id, bosqich=COMPLETE)
    await msg.answer(config.FOLLOWUP_FINISH_TEXT)
    lead = get_lead(msg.chat.id)
    ans = "\n".join(f"• {qlabel(q)}\n   <b>{lead['answers2'].get(q['key'], '-')}</b>"
                    for q in config.FOLLOWUP_QUESTIONS)
    await send_admins(msg.bot, f"📋 <b>Qo'shimcha ma'lumot to'liq</b>\n{lead_title(lead)}\n"
                               f"Xulosa: {lead['xulosa']}\n\n{ans}", docs_kb(lead["user_id"]))
    for key, fid in lead["photos"].items():
        caption = next((qlabel(q) for q in config.FOLLOWUP_QUESTIONS if q["key"] == key), key)
        for t in admin_targets():
            try:
                cap = f"{html.unescape(lead['name'])}: {caption}"[:1000]
                if fid.startswith("doc:"):
                    await msg.bot.send_document(t, fid[4:], caption=cap, parse_mode=None)
                else:
                    await msg.bot.send_photo(t, fid, caption=cap, parse_mode=None)
            except Exception as e:
                logging.warning("Rasm yuborilmadi %s: %s", t, e)


# ===================== BOSQICHLARNI ADMIN TOMONIDAN O'TKAZISH (Telegram va dashboard) =====================
# 1 · Ariza  →  2 · Taklif (send_offer)  →  3 · Hujjatlar (stage_to_docs)  →  ✅ Tasdiqlandi (stage_approve)

class _UserChat:
    """Foydalanuvchiga bot nomidan yozish uchun Message o'rnini bosadigan kichik obyekt."""
    def __init__(self, bot: Bot, uid: int):
        self.bot = bot
        self.chat = type("C", (), {"id": uid, "username": None})()

    async def answer(self, text, **kw):
        return await self.bot.send_message(self.chat.id, text, **kw)


def docs_kb(uid):
    return InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="✅ Hujjatlarni tasdiqlash", callback_data=f"ok:{uid}"),
        InlineKeyboardButton(text="↩️ Qayta so'rash", callback_data=f"re:{uid}")]])


async def _start_docs(bot: Bot, uid: int, intro: str):
    """Foydalanuvchini 3-bosqichga (hujjat savollari) o'tkazadi va birinchi savolni yuboradi."""
    m, st = _UserChat(bot, uid), user_state(bot, uid)
    try:
        await m.answer(intro, parse_mode=None)
    except Exception as e:
        logging.warning("Foydalanuvchiga yozib bo'lmadi %s: %s", uid, e)
        return False
    await st.clear()
    set_fields(uid, bosqich=AGREED)
    await st.update_data(answers2={}, photos={})
    await ask_followup(m, st, 0)
    return True


async def stage_to_docs(bot: Bot, lead):
    """2 → 3: odam roziligini boshqa yo'l bilan (masalan, telefonda) bergan bo'lsa, admin o'zi o'tkazadi."""
    return await _start_docs(bot, lead["user_id"],
                             "Arizangiz tasdiqlandi ✅ Endi rasmiylashtirish uchun bir nechta qo'shimcha ma'lumot kerak bo'ladi.")


async def stage_remind(bot: Bot, lead):
    """Javob bermayotgan odamga eslatma: taklifni qayta yuborish yoki joriy savolni qaytarish."""
    if lead["bosqich"] in (OFFER_SENT, ASKED, REFUSED):
        return await send_offer(bot, lead)
    if lead["bosqich"] == AGREED:
        m, st = _UserChat(bot, lead["user_id"]), user_state(bot, lead["user_id"])
        try:
            await m.answer("🔔 Eslatma: rasmiylashtirish uchun ma'lumotlarni yuborishni davom ettiring 👇", parse_mode=None)
        except Exception as e:
            logging.warning("Eslatma yuborilmadi %s: %s", lead["user_id"], e)
            return False
        await ask_followup(m, st, (await st.get_data()).get("fqi", 0))
        return True
    return False


async def stage_approve(bot: Bot, lead):
    """3 → ✅: hujjatlar tasdiqlandi — odamga xabar boradi."""
    try:
        await bot.send_message(lead["user_id"], config.APPROVED_TEXT, parse_mode=None)
    except Exception as e:
        logging.warning("Tasdiq xabari yuborilmadi %s: %s", lead["user_id"], e)
        return False
    await user_state(bot, lead["user_id"]).clear()
    set_fields(lead["user_id"], bosqich=APPROVED)
    return True


async def stage_reject(bot: Bot, lead, reason=""):
    """3 → 3: hujjatlarda kamchilik — odamdan qaytadan so'raladi."""
    why = f": {reason.strip()}" if reason and reason.strip() else ""
    return await _start_docs(bot, lead["user_id"],
                             f"Yuborgan ma'lumotlaringizda kamchilik bor{why}.\nIltimos, qaytadan yuboring 👇")


@dp.callback_query(F.data.startswith("ok:") | F.data.startswith("re:"))
async def admin_docs_decision(cb: CallbackQuery, bot: Bot):
    if not is_admin(cb.from_user.id):
        return await cb.answer("Faqat adminlar uchun", show_alert=True)
    lead = get_lead(int(cb.data[3:]))
    if not lead:
        return await cb.answer("Ariza topilmadi", show_alert=True)
    if lead["bosqich"] != COMPLETE:
        return await cb.answer(f"Hozirgi bosqich: {lead['bosqich'] or 'taklif yuborilmagan'}", show_alert=True)
    approve = cb.data.startswith("ok:")
    ok = await (stage_approve(bot, lead) if approve else stage_reject(bot, lead))
    if not ok:
        return await cb.answer("Yuborib bo'lmadi (botni bloklagan bo'lishi mumkin)", show_alert=True)
    await cb.answer("Bajarildi ✅")
    label = "✅ Tasdiqlandi" if approve else "↩️ Qayta so'raldi"
    await cb.message.edit_reply_markup(reply_markup=InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text=f"{label} ({cb.from_user.first_name})", callback_data="noop")]]))

# ===================== ADMIN =====================

@dp.message(Command("myid"))
async def myid(msg: Message):
    await msg.answer(f"Sizning ID: <code>{msg.from_user.id}</code>\nChat ID: <code>{msg.chat.id}</code>")


@dp.message(Command("help"))
async def help_cmd(msg: Message):
    if not is_admin(msg.from_user.id):
        return
    await msg.answer(
        "<b>Admin buyruqlari</b>\n"
        "/link post1 — e'lon uchun maxsus havola (manba kuzatiladi)\n"
        "/stats — umumiy statistika\n"
        "/export — Excel (viloyatlar bo'yicha varaqlarga saralangan)\n"
        "/leads maqul | korib | issiq | tugallanmagan | <i>viloyat yoki tuman</i> — ro'yxat\n"
        "/maqul — ✅ Ma'qul lidlar va «Davom ettirish» (taklif yuborish) tugmalari\n"
        "/rasmlar — barcha rasmlar ZIP faylda (/rasmlar <i>ism yoki tuman</i> — bittasiniki)\n"
        "/recheck — barcha lidlarni filiallar bilan qayta solishtirish\n"
        "/filiallar — filiallar bazasi holati\n"
        "/setgroup — shu guruhga yangi lidlar keladi (guruhda yozing)\n"
        "/unsetgroup — xabarlar yana adminlarga shaxsan keladi")


@dp.message(Command("link"))
async def link(msg: Message, command: CommandObject, bot: Bot):
    if not is_admin(msg.from_user.id):
        return
    src = re.sub(r"[^A-Za-z0-9_-]", "", command.args or "")[:64] or "kanal"
    me = await bot.get_me()
    await msg.answer(f"E'longa qo'yiladigan havola:\n<code>https://t.me/{me.username}?start={src}</code>\n\n"
                     f"Shu havoladan kelgan lidlar «{src}» manbasi bilan belgilanadi.")


@dp.message(Command("setgroup"))
async def setgroup(msg: Message):
    if not is_admin(msg.from_user.id):
        return
    set_setting("group", msg.chat.id)
    await msg.answer("✅ Endi yangi lidlar shu chatga yuboriladi.")


@dp.message(Command("unsetgroup"))
async def unsetgroup(msg: Message):
    if not is_admin(msg.from_user.id):
        return
    with db() as con:
        con.execute("DELETE FROM settings WHERE k='group'")
    await msg.answer("✅ Lidlar endi adminlarga shaxsan yuboriladi.")


def counter_block(title, cnt):
    lines = "\n".join(f"  {k}: <b>{v}</b>" for k, v in cnt.most_common())
    return f"<b>{title}</b>\n{lines or '  —'}"


@dp.message(Command("stats"))
async def stats(msg: Message):
    if not is_admin(msg.from_user.id):
        return
    leads = all_leads()
    parts = [f"📊 Jami lidlar: <b>{len(leads)}</b>",
             counter_block("Xulosa", Counter(l["xulosa"] for l in leads if l["holat"] == DONE)),
             counter_block("Taklif bosqichi", Counter(l["bosqich"] for l in leads if l["bosqich"])),
             counter_block("Holat", Counter(l["holat"] for l in leads)),
             counter_block("Daraja", Counter(l["level"] for l in leads)),
             counter_block("Viloyat", Counter(l["viloyat"] for l in leads)),
             counter_block("Tuman", Counter(f"{l['viloyat']}, {l['tuman']}" for l in leads)),
             counter_block("Manba", Counter(l["source"] for l in leads))]
    for q in config.QUESTIONS:
        if q.get("type") == "choice":
            parts.append(counter_block(qlabel(q), Counter(l["answers"].get(q["key"], "-") for l in leads)))
    await msg.answer("\n\n".join(parts)[:4096])


@dp.message(Command("leads"))
async def leads_cmd(msg: Message, command: CommandObject):
    if not is_admin(msg.from_user.id):
        return
    f = (command.args or "").lower().strip()
    leads = all_leads()
    marks = {"maqul": "✅", "ma'qul": "✅", "korib": "⚠", "ko'rib": "⚠", "maqul emas": "❌", "ma'qul emas": "❌"}
    if f in marks:
        leads = [l for l in leads if l["xulosa"].startswith(marks[f])]
    elif f:
        leads = [l for l in leads if f in l["level"].lower() or f in l["holat"].lower() or f in l["xulosa"].lower().replace("'", "") or f in l["viloyat"].lower()
                 or f in l["tuman"].lower() or f in l["source"].lower()
                 or any(f in str(v).lower() for v in l["answers"].values())]
    if not leads:
        return await msg.answer("Lid topilmadi.")
    lines = [f"{i}. {l['xulosa'][:1]}{l['level'][:2]} {l['name']} — {l['phone']} — {l['viloyat']}, {l['tuman']}"
             for i, l in enumerate(leads[:50], 1)]
    more = f"\n\n…va yana {len(leads) - 50} ta. To'liq ro'yxat: /export" if len(leads) > 50 else ""
    await msg.answer(f"Topildi: <b>{len(leads)}</b>\n\n" + "\n".join(lines) + more)


def build_excel(leads):
    wb = Workbook()
    head = ["Xulosa", "Sabab", "Eng yaqin filial", "Masofa, km", "Tumandagi filiallar", "Aholi zichligi",
            "Status", "Izoh", "Keyingi aloqa", "Holat", "Daraja", "Ball", "Ism", "Telefon", "Username", "Viloyat", "Tuman", "Xarita", "Manba"] + \
           [qlabel(q) for q in config.QUESTIONS] + ["Bosqich"] + \
           [qlabel(q) for q in config.FOLLOWUP_QUESTIONS] + ["Sana", "Yangilangan"]
    fills = {"🔥": "FFC7CE", "🌤": "FFEB9C", "❄️": "DDEBF7"}

    photo_qs = [q for q in config.FOLLOWUP_QUESTIONS if q["type"] == "photo"]

    def add_photos(ws, rows):
        """Rasmlarni katakka joylaydi (asl sifatda, kichraytirib ko'rsatiladi)."""
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.utils import get_column_letter
        for q in photo_qs:
            ws.column_dimensions[get_column_letter(head.index(qlabel(q)) + 1)].width = 28
        for r, l in enumerate(rows, start=2):
            placed = False
            for q in photo_qs:
                path = photo_path(l["user_id"], q["key"])
                if not path.exists():
                    continue
                img = XLImage(str(path))
                scale = 150 / img.height
                img.height, img.width = 150, int(img.width * scale)
                col = head.index(qlabel(q)) + 1
                ws.cell(r, col).value = None
                ws.add_image(img, f"{get_column_letter(col)}{r}")
                placed = True
            if placed:
                ws.row_dimensions[r].height = 118

    def fill_sheet(ws, rows):
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for l in rows:
            ws.append([l["xulosa"], l["sabab"] or "", l["eng_yaqin"] or "",
                       float(l["masofa_km"]) if l["masofa_km"] else None,
                       int(l["tumanda_filial"]) if l["tumanda_filial"] is not None else None, l["zichlik"] or "", l["status"], l["izoh"], l["aloqa_sana"], l["holat"], l["level"], l["score"], html.unescape(l["name"]), l["phone"],
                       ("@" + l["username"]) if l["username"] else "", l["viloyat"], l["tuman"],
                       map_link(l), l["source"]] +
                      [html.unescape(l["answers"].get(q["key"], "")) for q in config.QUESTIONS] +
                      [l["bosqich"]] +
                      [html.unescape(l["answers2"].get(q["key"], "")) for q in config.FOLLOWUP_QUESTIONS] +
                      [l["created"], l["updated"]])
            color = next((v for k, v in fills.items() if l["level"].startswith(k)), None)
            if color:
                ws.cell(ws.max_row, head.index("Daraja") + 1).fill = PatternFill("solid", fgColor=color)
            vcolor = {"✅": "C6EFCE", "⚠": "FFEB9C", "❌": "FFC7CE"}.get(l["xulosa"][:1])
            if vcolor:
                ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=vcolor)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(
                12, min(40, max(len(str(c.value or "")) for c in col[1:] or col) + 2))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Barchasi"
    fill_sheet(ws, leads)
    taklif = [l for l in leads if l["bosqich"]]
    if taklif:
        ws_t = wb.create_sheet("Taklif jarayoni")
        fill_sheet(ws_t, taklif)
        add_photos(ws_t, taklif)
    for title, mark in (("Ma'qul", "✅"), ("Ko'rib chiqish", "⚠"), ("Ma'qul emas", "❌")):
        rows = [l for l in leads if l["holat"] == DONE and l["xulosa"].startswith(mark)]
        if rows:
            fill_sheet(wb.create_sheet(title), rows)
    drafts = [l for l in leads if l["holat"] == DRAFT]
    if drafts:
        fill_sheet(wb.create_sheet("Tugallanmagan"), drafts)
    for region in REGIONS:
        rows = [l for l in leads if l["viloyat"] == region]
        if rows:
            fill_sheet(wb.create_sheet(re.sub(r"[\[\]:*?/\\']", "", region)[:31]), rows)

    st = wb.create_sheet("Statistika")
    groups = [("Xulosa", lambda l: l["xulosa"]),
              ("Holat", lambda l: l["holat"]),
              ("Daraja", lambda l: l["level"]),
              ("Viloyat", lambda l: l["viloyat"]),
              ("Tuman", lambda l: f"{l['viloyat']}, {l['tuman']}"),
              ("Manba", lambda l: l["source"])] + \
             [(qlabel(q), lambda l, k=q["key"]: l["answers"].get(k, "-"))
              for q in config.QUESTIONS if q.get("type") == "choice"]
    for title, key in groups:
        st.append([title])
        st.cell(st.max_row, 1).font = Font(bold=True)
        for k, v in Counter(key(l) for l in leads).most_common():
            st.append([k, v])
        st.append([])
    st.column_dimensions["A"].width = 50

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@dp.message(Command("maqul"))
async def maqul_cmd(msg: Message):
    if not is_admin(msg.from_user.id):
        return
    leads = [l for l in all_leads() if l["holat"] == DONE and l["xulosa"].startswith("✅")]
    if not leads:
        return await msg.answer("Hozircha ✅ Ma'qul lid yo'q.")
    new = [l for l in leads if not l["bosqich"]]
    lines = [f"{i}. {(l['bosqich'] or '🆕')[:1]} {l['name']} — {l['tuman']} — {l['phone']}"
             for i, l in enumerate(leads, 1)]
    text = (f"✅ <b>Ma'qul lidlar: {len(leads)}</b>\n\n" + "\n".join(lines) +
            "\n\n<i>🆕 taklif yuborilmagan · 📨 yuborildi · ❓ savol berdi · 🤝 roziman · "
            "📋 ma'lumot to'liq · 🚫 rad etdi</i>")
    if new:
        text += f"\n\n📨 Taklif yuborish uchun pastdagi tugmani bosing ({len(new)} ta)."
    rows = [[InlineKeyboardButton(text=f"📨 {html.unescape(l['name'])} — {l['tuman']}"[:60],
                                  callback_data=f"go:{l['user_id']}")] for l in new[:40]]
    if len(new) > 1:
        rows.append([InlineKeyboardButton(text=f"📨 Barchasiga yuborish ({len(new)})", callback_data="goall")])
    await msg.answer(text[:4096], reply_markup=InlineKeyboardMarkup(inline_keyboard=rows) if rows else None)


@dp.message(Command("recheck"))
async def recheck(msg: Message):
    if not is_admin(msg.from_user.id):
        return
    leads = [l for l in all_leads() if l["holat"] == DONE]
    await msg.answer(f"⏳ {len(leads)} ta lid qayta tekshirilmoqda...")
    for l in leads:
        await run_evaluation(l)
        await asyncio.sleep(1)
    c = Counter(l["xulosa"] for l in all_leads() if l["holat"] == DONE)
    await msg.answer("✅ Tayyor.\n" + "\n".join(f"{k}: <b>{v}</b>" for k, v in c.most_common()))


@dp.message(Command("filiallar"))
async def filiallar_cmd(msg: Message):
    if not is_admin(msg.from_user.id):
        return
    try:
        allb = branches.load_branches()
    except Exception as e:
        return await msg.answer(f"❌ Filiallar fayli o'qilmadi: {e}")
    act = branches.active_branches()
    await msg.answer(f"🏢 Filiallar bazasi: <b>{len(allb)}</b> ta, aktiv va koordinatali: <b>{len(act)}</b> ta\n"
                     f"Radius: shahar {config.SHAHAR_RADIUS_KM} km, tuman {config.TUMAN_RADIUS_KM} km")


@dp.message(Command("export"))
async def export(msg: Message):
    if not is_admin(msg.from_user.id):
        return
    leads = all_leads()
    if not leads:
        return await msg.answer("Hozircha lid yo'q.")
    await ensure_photos(msg.bot, leads)
    name = f"lidlar_{datetime.now():%Y%m%d_%H%M}.xlsx"
    await msg.answer_document(BufferedInputFile(build_excel(leads), filename=name),
                              caption=f"Jami: {len(leads)} ta lid "
                                      f"({sum(l['holat'] == DONE for l in leads)} ta to'liq, "
                                      f"{sum(l['holat'] == DRAFT for l in leads)} ta tugallanmagan)\n"
                                      f"Rasmlar «Taklif jarayoni» varag'ida. Hammasini alohida fayl qilib "
                                      f"olish: /rasmlar")


def safe_name(text):
    return re.sub(r'[\\/:*?"<>|\']+', "", html.unescape(text)).strip()[:60] or "lid"


@dp.message(Command("rasmlar"))
async def rasmlar_cmd(msg: Message, command: CommandObject):
    """Barcha (yoki qidiruvga mos) lidlarning rasmlarini ZIP qilib yuboradi."""
    if not is_admin(msg.from_user.id):
        return
    import zipfile
    f = (command.args or "").lower().strip()
    leads = [l for l in all_leads() if l["photos"]]
    if f:
        leads = [l for l in leads if f in l["name"].lower() or f in l["tuman"].lower()
                 or f in l["viloyat"].lower() or f in l["phone"]]
    if not leads:
        return await msg.answer("Rasm yuborgan lid topilmadi.")
    await ensure_photos(msg.bot, leads)
    labels = {q["key"]: q.get("name") or q["key"] for q in config.FOLLOWUP_QUESTIONS}
    buf, count = io.BytesIO(), 0
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for l in leads:
            folder = safe_name(f"{l['name']} - {l['tuman']} - {l['phone']}")
            for key in l["photos"]:
                path = photo_path(l["user_id"], key)
                if path.exists():
                    z.write(path, f"{folder}/{labels.get(key, key)}.jpg")
                    count += 1
    if not count:
        return await msg.answer("Rasmlarni yuklab bo'lmadi.")
    await msg.answer_document(BufferedInputFile(buf.getvalue(), filename=f"rasmlar_{datetime.now():%Y%m%d_%H%M}.zip"),
                              caption=f"{len(leads)} ta lid, {count} ta rasm")

# ===================== ISHGA TUSHIRISH =====================

async def main():
    if not config.BOT_TOKEN:
        raise SystemExit("config.py ichida BOT_TOKEN ni kiriting!")
    import socket
    import sys
    from logging.handlers import RotatingFileHandler
    (BASE / "logs").mkdir(exist_ok=True)
    handlers = [RotatingFileHandler(BASE / "logs" / "bot.log", maxBytes=2_000_000, backupCount=3, encoding="utf-8")]
    if sys.stderr:   # konsoldan ishga tushirilgan bo'lsa, ekranga ham yozamiz (pythonw da stderr yo'q)
        handlers.append(logging.StreamHandler())
    logging.basicConfig(level=logging.INFO, handlers=handlers,
                        format="%(asctime)s %(levelname)s %(name)s: %(message)s")
    # Bitta nusxa: dashboard porti band bo'lsa — bot allaqachon ishlayapti
    with socket.socket() as sck:
        if sck.connect_ex(("127.0.0.1", config.DASHBOARD_PORT)) == 0:
            logging.warning("Bot allaqachon ishlayapti (%s-port band) — ikkinchi nusxa ishga tushmaydi", config.DASHBOARD_PORT)
            return
    init_db()
    bot = Bot(config.BOT_TOKEN, default=DefaultBotProperties(parse_mode="HTML"))
    import dashboard
    try:
        await dashboard.start(bot, sys.modules[__name__])
    except OSError as e:
        logging.error("Dashboard ishga tushmadi (%s-port band bo'lishi mumkin): %s", config.DASHBOARD_PORT, e)
    print("Bot ishga tushdi. To'xtatish: Ctrl+C")
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
